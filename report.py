"""Excel report for the CEO. First sheet = your standard monthly format; extra sheets
back it up (daily register, day-wise detail, permissions & overtime, review items, rules used)."""
from __future__ import annotations

import calendar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from engine import fmt_min, STATUS_LABEL

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name="Calibri", size=11, bold=True)
NORM = Font(name="Calibri", size=11)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
GREY = PatternFill("solid", fgColor="D9D9D9")
HEAD_EXTRA = PatternFill("solid", fgColor="DDEBF7")

DAYS_FMT = '[=0]"-";[=1]0" Day";General" Days"'
FILL = {
    "P": "C6EFCE", "OD": "B7E1E8", "WFH": "D9D2E9", "HPL": "FCE4D6", "HLOP": "F8CBAD",
    "PL": "FFF2CC", "LOP": "FFC7CE", "WO": "E7E6E6", "H": "BDD7EE", "UNMARKED": "FFFF00", "PENDING": "F2F2F2",
}
ACTION = {
    "ABSENT_LOP": "Absent = LOP. If the leave was intimated/approved, change the code to PL",
    "UNMARKED": "Decide: Present / Paid leave / LOP / OD / WFH and save the day",
    "MISSING_OUT": "Enter the OUT time from the machine / manager, or mark half day",
    "MISSING_IN": "Enter the IN time",
    "NO_PUNCH": "Enter punches, or change the code to OD / WFH / leave",
    "UNKNOWN_CODE": "Replace with a valid code (P, A, WO, H, PL, LOP, OD, WFH, HPL, HLOP)",
    "BAD_TIME": "Correct the time (HH:MM)",
    "BAD_TIMES": "Correct the IN / OUT times",
    "BAD_PERM": "Permission hours must be a number",
    "EARLY_EXIT": "Check approval: permission / half day / normal?",
    "SHORT_HOURS": "Check whether this is a half day (HPL / HLOP)",
    "PUNCH_ON_NON_WORKING": "Check: comp-off / overtime, or wrong code",
    "PERM_OVER_QUOTA": "Deduction applied per the rule in Settings - confirm with management",
    "OCR_VERIFY": "Verify against the photo, then save the row",
    "CODE_INFERRED": "Information only",
}


def _month_title(year, month, st):
    return f"Attendance for {calendar.month_name[month]} Month-{str(year)[2:]}"


def _check(s):
    if not s["balanced"]:
        return "CHECK TOTALS"
    if s["unresolved"]:
        return f"{s['unresolved']} day(s) unmarked"
    if s["pending"]:
        return f"{s['pending']} day(s) pending"
    return "OK"


def _hdr(ws, row, headers, fills=None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.alignment, c.border = BOLD, CENTER, BORDER
        if fills and fills.get(i):
            c.fill = fills[i]


def build_workbook(results, year, month, st, path):
    wb = Workbook()

    # ============ 1. standard summary ============
    ws = wb.active
    ws.title = "Attendance Summary"
    heads = ["S.No", "No.Of Employees ", "Designation ", "No.Of Days ", "No.Of Working ", "No.Of Weekoff ",
             "Granted Holidays", "No.Of present", "Permitted Leave ", "Loss Of Pay ",
             "Payable Days (Salary)", "Permission Hrs Used (of quota)", "OD Days", "WFH Days", "OT Hours", "Check"]
    ncols = len(heads)
    for r, text in ((1, st["company_name"]), (2, _month_title(year, month, st))):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=text)
        c.font, c.alignment = BOLD, CENTER
        for col in range(1, ncols + 1):
            ws.cell(row=r, column=col).border = BORDER
    _hdr(ws, 3, heads, {i: HEAD_EXTRA for i in range(11, ncols + 1)})
    ws.row_dimensions[3].height = 32
    quota = st["monthly_permission_hours"]
    for n, res in enumerate(results, start=1):
        e, s = res["emp"], res["summary"]
        row = 3 + n
        vals = [n, e["name"], e["designation"], s["days"], s["working"], s["weekoff"], s["holidays"],
                s["present"], s["leave"], s["lop"], s["payable"],
                f"{s['perm_used']:g} / {quota:g}", s["od"], s["wfh"],
                fmt_min(s["ot_min"]) if s["ot_min"] else "-",
                _check(s)]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.font, c.border = NORM, BORDER
            c.alignment = LEFT if col in (2, 3) else CENTER
            if 4 <= col <= 10 or col in (11, 13, 14):
                c.number_format = DAYS_FMT
        if _check(s) != "OK":
            ws.cell(row=row, column=ncols).fill = PatternFill("solid", fgColor="FFFF00")
    widths = [5, 22, 30, 11, 13, 13, 14, 13, 14, 12, 14, 18, 9, 9, 10, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    note = 3 + len(results) + 2
    notes = ["Present = Present + Official On Duty + Work From Home (half days count 0.5).",
             "Working days = Days in month - Week off - Granted holidays.  Present + Permitted leave + Loss of pay = Working days.",
             "Payable days = Days in month - Loss of pay (week offs and granted holidays are paid).",
             f"Late arrival after {st['late_grace_until']} = 1 hour permission; {quota:g} hours of permission allowed per month."]
    for i, t in enumerate(notes):
        ws.cell(row=note + i, column=1, value=t).font = Font(name="Calibri", size=9, italic=True)
    ws.freeze_panes = "D4"

    # ============ 2. daily register ============
    wr = wb.create_sheet("Daily Register")
    ndays = calendar.monthrange(year, month)[1]
    wr.cell(row=1, column=1, value=f"{st['company_name']} - Daily register {calendar.month_name[month]} {year}").font = BOLD
    wr.cell(row=3, column=1, value="Employee").font = BOLD
    for d in range(1, ndays + 1):
        a = wr.cell(row=2, column=1 + d, value=calendar.day_abbr[calendar.weekday(year, month, d)][:2])
        b = wr.cell(row=3, column=1 + d, value=d)
        for c in (a, b):
            c.font, c.alignment, c.border = BOLD, CENTER, BORDER
        wr.column_dimensions[get_column_letter(1 + d)].width = 5.2
    for k, lab in enumerate(["Present", "Leave", "LOP"]):
        c = wr.cell(row=3, column=2 + ndays + k, value=lab)
        c.font, c.alignment, c.border = BOLD, CENTER, BORDER
        wr.column_dimensions[get_column_letter(2 + ndays + k)].width = 9
    wr.column_dimensions["A"].width = 22
    for n, res in enumerate(results):
        row = 4 + n
        c = wr.cell(row=row, column=1, value=res["emp"]["name"])
        c.font, c.border = NORM, BORDER
        for dr in res["days"]:
            text = {"UNMARKED": "?", "PENDING": "-"}.get(dr["status"], dr["status"])
            if dr["perm_hours"] > 0 and dr["status"] == "P":
                text += "*"
            c = wr.cell(row=row, column=1 + dr["day"], value=text)
            c.font, c.alignment, c.border = Font(name="Calibri", size=9), CENTER, BORDER
            c.fill = PatternFill("solid", fgColor=FILL.get(dr["status"], "FFFFFF"))
        s = res["summary"]
        for k, v in enumerate([s["present"], s["leave"], s["lop"]]):
            c = wr.cell(row=row, column=2 + ndays + k, value=v)
            c.font, c.alignment, c.border = NORM, CENTER, BORDER
    leg = 5 + len(results)
    wr.cell(row=leg, column=1, value="Legend").font = BOLD
    for i, (k, lab) in enumerate(STATUS_LABEL.items()):
        c = wr.cell(row=leg + 1 + i, column=1, value={"UNMARKED": "?", "PENDING": "-"}.get(k, k))
        c.fill, c.alignment, c.border = PatternFill("solid", fgColor=FILL[k]), CENTER, BORDER
        wr.cell(row=leg + 1 + i, column=2, value=lab)
    wr.cell(row=leg + 2 + len(STATUS_LABEL), column=1, value="P*").alignment = CENTER
    wr.cell(row=leg + 2 + len(STATUS_LABEL), column=2, value="Present, permission hours used that day")
    wr.freeze_panes = "B4"

    # ============ 3. day-wise detail ============
    wd = wb.create_sheet("Day-wise Detail")
    dh = ["Employee", "Date", "Day", "Sheet code", "Final status", "IN", "OUT", "Worked (hh:mm)",
          "Overtime (hh:mm)", "Permission hrs", "Permission running total", "Deduction (days)", "Remarks / note", "Source", "Flags"]
    _hdr(wd, 1, dh)
    r = 2
    for res in results:
        for dr in res["days"]:
            vals = [res["emp"]["name"], dr["date"], dr["weekday"], dr["raw_code"], dr["status"], dr["in"], dr["out"],
                    fmt_min(dr["worked_min"]) if dr["worked_min"] is not None else "",
                    fmt_min(dr["ot_min"]) if dr["ot_min"] else "",
                    dr["perm_hours"] or "", dr["perm_running"] if dr["perm_hours"] else "",
                    dr["deduction"] or "", (dr["note"] + " " + dr["remarks"]).strip(), dr["source"],
                    "; ".join(f[2] for f in dr["flags"])]
            for col, v in enumerate(vals, start=1):
                c = wd.cell(row=r, column=col, value=v)
                c.font, c.border = NORM, BORDER
                c.alignment = LEFT if col in (1, 13, 15) else CENTER
            wd.cell(row=r, column=5).fill = PatternFill("solid", fgColor=FILL.get(dr["status"], "FFFFFF"))
            r += 1
    for i, w in enumerate([22, 12, 6, 14, 12, 8, 8, 12, 12, 11, 14, 11, 30, 9, 70], start=1):
        wd.column_dimensions[get_column_letter(i)].width = w
    wd.auto_filter.ref = f"A1:{get_column_letter(len(dh))}{max(r - 1, 2)}"
    wd.freeze_panes = "C2"

    # ============ 4. permissions & overtime ============
    wp = wb.create_sheet("Permissions & OT")
    _hdr(wp, 1, ["Employee", "Date", "IN time", "Permission hrs", "Running total", f"Quota {quota:g} hrs", "Deduction (days)", "Reason"])
    r = 2
    for res in results:
        for dr in res["days"]:
            if dr["perm_hours"] > 0:
                reason = (f"Late arrival ({dr['late_min']} min after {st['shift_start']})" if dr["late_min"] and dr["in"] > st["late_grace_until"]
                          else "Permission entered / early exit")
                vals = [res["emp"]["name"], dr["date"], dr["in"], dr["perm_hours"], dr["perm_running"],
                        "OVER" if dr["perm_over"] else "within", dr["deduction"] or "", reason]
                for col, v in enumerate(vals, start=1):
                    c = wp.cell(row=r, column=col, value=v)
                    c.font, c.border, c.alignment = NORM, BORDER, CENTER if col > 1 else LEFT
                r += 1
    if r == 2:
        wp.cell(row=2, column=1, value="No permissions used this month.").font = NORM
    r += 2
    _hdr(wp, r, ["Employee", "Permission used (hrs)", "Quota left (hrs)", "Over quota (hrs)", "Overtime (hh:mm)"])
    for res in results:
        r += 1
        s = res["summary"]
        for col, v in enumerate([res["emp"]["name"], s["perm_used"], s["perm_left"], s["perm_excess"],
                                 fmt_min(s["ot_min"]) if s["ot_min"] else "-"], start=1):
            c = wp.cell(row=r, column=col, value=v)
            c.font, c.border, c.alignment = NORM, BORDER, CENTER if col > 1 else LEFT
    for i, w in enumerate([24, 14, 12, 15, 14, 14, 15, 44], start=1):
        wp.column_dimensions[get_column_letter(i)].width = w

    # ============ 5. review items ============
    wv = wb.create_sheet("Review Items")
    _hdr(wv, 1, ["Level", "Employee", "Date", "Issue", "Suggested action"])
    r = 2
    for res in results:
        for dr in res["days"]:
            for lvl, key, msg in dr["flags"]:
                if lvl == "info":
                    continue
                for col, v in enumerate([lvl.upper(), res["emp"]["name"], dr["date"], msg, ACTION.get(key, "")], start=1):
                    c = wv.cell(row=r, column=col, value=v)
                    c.font, c.border, c.alignment = NORM, BORDER, LEFT
                wv.cell(row=r, column=1).fill = PatternFill("solid", fgColor="FFC7CE" if lvl == "error" else "FFEB9C")
                r += 1
    if r == 2:
        wv.cell(row=2, column=1, value="Nothing to review.").font = NORM
    for i, w in enumerate([9, 22, 12, 60, 60], start=1):
        wv.column_dimensions[get_column_letter(i)].width = w
    wv.auto_filter.ref = f"A1:E{max(r - 1, 2)}"

    # ============ 6. rules used ============
    wu = wb.create_sheet("Rules Used")
    rules = [
        ("Shift", f"{st['shift_start']} - {st['shift_end']}  (paid hours {st['paid_hours']})"),
        ("Late arrival", f"After {st['late_grace_until']} = 1 hour permission (+1 hour for each further hour late)"),
        ("Monthly permission quota", f"{st['monthly_permission_hours']:g} hours"),
        ("Beyond the quota", {"half_day_each": "0.5 day LOP for each instance beyond the quota",
                              "proportional": "Hours beyond the quota / paid hours = LOP days",
                              "none": "No deduction (reported only)"}.get(st["excess_permission_rule"], st["excess_permission_rule"])),
        ("Sheet code A (absent)", f"Counted as {st['absent_default']} (set PL for intimated leave)"),
        ("Early exit as permission", "Yes" if st["early_exit_as_permission"] else "No (flagged for review)"),
        ("Overtime", f"Worked hours (OUT - IN) beyond {st['paid_hours']}"),
        ("Week off days", ", ".join(calendar.day_name[d] for d in st["weekly_off_days"]) + " (used only when a day has no code)"),
    ]
    _hdr(wu, 1, ["Rule", "Setting used for this report"])
    for i, (a, b) in enumerate(rules, start=2):
        for col, v in enumerate((a, b), start=1):
            c = wu.cell(row=i, column=col, value=v)
            c.font, c.border, c.alignment = NORM, BORDER, LEFT
    wu.column_dimensions["A"].width = 28
    wu.column_dimensions["B"].width = 90

    for sheet in wb.worksheets:                      # print / PDF friendly: landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
    wd.print_title_rows = "1:1"
    wb.save(path)
    return path
